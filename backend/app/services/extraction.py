from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".xlsx", ".xls"}
MAX_FILE_MB = 100
MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix in {".xlsx", ".xls"}:
        return _extract_xlsx(path)
    raise ValueError(f"Formato no soportado: {suffix}")


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Página {index}]\n{text}")
    return "\n\n".join(pages).strip()


def _extract_docx(path: Path) -> str:
    document = DocxDocument(str(path))
    parts = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_xlsx(path: Path) -> str:
    workbook = load_workbook(str(path), data_only=True, read_only=True)
    chunks = []
    for sheet in workbook.worksheets:
        chunks.append(f"[Hoja {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell).strip() for cell in row if cell not in (None, "")]
            if values:
                chunks.append(" | ".join(values))
    workbook.close()
    return "\n".join(chunks).strip()
