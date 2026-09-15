from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..compliance import DROPDOWN, display_compliance
from ..models import Deliverable, Expediente, ScopeItem

HEADER_FILL = "1B365D"
STRIPE = "F2F5F9"
WHITE = "FFFFFF"
BORDER = "D3D3D3"
GREEN_FILL = "D4EDDA"
GREEN_FONT = "155724"
YELLOW_FILL = "FFF3CD"
YELLOW_FONT = "856404"
RED_FILL = "F8D7DA"
RED_FONT = "721C24"

CHECKLIST_HEADERS = [
    "A. Numeral del ítem",
    "B. Bases integradas / Documento base",
    "C. Consultas",
    "D. Propuesta",
    "E. Contrato",
    "F. Seguimiento Operativo",
]
DELIVERABLE_HEADERS = [
    "A. Numeral",
    "B. Entregable identificado",
    "C. Plazo exigido",
    "D. Referencia documental",
    "E. Seguimiento Operativo",
]


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _border() -> Border:
    thin = Side(style="thin", color=BORDER)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _header_row(ws, values: list[str]) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=1, column=col, value=value)
        cell.fill = _fill(HEADER_FILL)
        cell.font = Font(name="Arial", bold=True, color=WHITE, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _border()
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = True


def _write_body(ws, rows: list[list], center_cols: set[int], status_col: int, widths: list[float]) -> None:
    last = max(len(rows) + 1, 2)
    for excel_row, values in enumerate(rows, start=2):
        stripe = _fill(STRIPE) if excel_row % 2 == 0 else _fill(WHITE)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=excel_row, column=col, value=value if value else None)
            cell.fill = stripe
            cell.font = Font(name="Arial", size=10)
            cell.border = _border()
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
                horizontal="center" if col in center_cols else "left",
            )
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    end_col = get_column_letter(len(widths))
    ws.auto_filter.ref = f"A1:{end_col}{last}"
    letter = get_column_letter(status_col)
    status_range = f"{letter}2:{letter}{last}"
    dv = DataValidation(
        type="list",
        formula1=DROPDOWN,
        allow_blank=False,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Valor no permitido",
        error="Seleccione ☐ Pendiente, ☑ Cumplido o ☒ Incumplido.",
        promptTitle="Seguimiento Operativo",
        prompt="☐ Pendiente / ☑ Cumplido / ☒ Incumplido",
        showInputMessage=True,
    )
    ws.add_data_validation(dv)
    dv.add(status_range)
    rules = [
        ("☑ Cumplido", GREEN_FILL, GREEN_FONT),
        ("☐ Pendiente", YELLOW_FILL, YELLOW_FONT),
        ("☒ Incumplido", RED_FILL, RED_FONT),
    ]
    for label, fill, font in rules:
        ws.conditional_formatting.add(
            status_range,
            CellIsRule(
                operator="equal",
                formula=[f'"{label}"'],
                fill=_fill(fill),
                font=Font(name="Arial", size=10, bold=True, color=font),
                stopIfTrue=True,
            ),
        )
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True


async def build_workbook(expediente: Expediente) -> bytes:
    items = await ScopeItem.find(ScopeItem.expediente_id == expediente.id).sort("+sort_order").to_list()
    deliverables = await Deliverable.find(Deliverable.expediente_id == expediente.id).sort("+sort_order").to_list()

    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist_Alcance"
    _header_row(ws, CHECKLIST_HEADERS)
    checklist_rows = [
        [
            item.numeral,
            item.bases,
            item.consultas,
            item.propuesta,
            item.contrato,
            display_compliance(item.compliance),
        ]
        for item in items
    ]
    _write_body(ws, checklist_rows, center_cols={1, 6}, status_col=6, widths=[22, 45, 35, 35, 45, 20])

    ws2 = wb.create_sheet("Entregables")
    _header_row(ws2, DELIVERABLE_HEADERS)
    deliverable_rows = [
        [
            item.numeral,
            item.name,
            item.due_term,
            item.reference,
            display_compliance(getattr(item, "compliance", None)),
        ]
        for item in deliverables
    ]
    _write_body(ws2, deliverable_rows, center_cols={1, 5}, status_col=5, widths=[22, 50, 35, 35, 20])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
